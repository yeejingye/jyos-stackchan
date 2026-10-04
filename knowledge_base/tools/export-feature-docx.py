"""Generate a Word view from the Markdown token payload prepared by export-feature.mjs."""
import json
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

data = json.loads(Path(sys.argv[1]).read_text())
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.75)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt({"Normal": 10.5, "Title": 27, "Subtitle": 10, "Heading 1": 16, "Heading 2": 12, "Heading 3": 11}[name])
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.12
    if name.startswith("Heading"):
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles["Normal"].paragraph_format.widow_control = True
doc.core_properties.title = data["title"]
doc.core_properties.subject = "Derived feature specification view"
doc.core_properties.author = "Codex"
doc.core_properties.comments = f"Source SHA256 {data['hash']}; generated {data['generated']}"


def shade(cell_or_paragraph, fill):
    props = cell_or_paragraph._tc.get_or_add_tcPr() if hasattr(cell_or_paragraph, "_tc") else cell_or_paragraph._p.get_or_add_pPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def hyperlink(paragraph, label, target):
    if not urlparse(target).scheme and not target.startswith("#"):
        target = (Path(data["source"]).parent / target).resolve().as_uri()
    rel = paragraph.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "216B68")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, tokens, bold=False, italic=False):
    for token in tokens:
        kind = token["type"]
        if kind == "link":
            hyperlink(paragraph, token["text"], token["href"])
        elif kind in ("strong", "em", "text") and token.get("tokens"):
            inline(paragraph, token["tokens"], bold or kind == "strong", italic or kind == "em")
        elif kind == "br":
            paragraph.add_run().add_break()
        else:
            run = paragraph.add_run(token.get("text", token.get("raw", "")))
            run.bold, run.italic = bold, italic
            if kind == "codespan":
                run.font.name = "Liberation Mono"
                run.font.size = Pt(9)


def block(token, list_style=None):
    kind = token["type"]
    if kind == "heading":
        doc.add_paragraph(token["text"], "Title" if token["depth"] == 1 else f"Heading {min(token['depth'] - 1, 3)}")
        if token["depth"] == 1:
            date = datetime.fromisoformat(data["generated"].replace("Z", "+00:00")).strftime("%d %B %Y")
            doc.add_paragraph(f"Feature specification | {date} | Source SHA256 {data['hash'][:12]}", "Subtitle")
    elif kind in ("paragraph", "text"):
        paragraph = doc.add_paragraph(style=list_style)
        inline(paragraph, token.get("tokens", [{"type": "text", "text": token.get("text", "")}]))
    elif kind == "list":
        for item in token["items"]:
            for child in item.get("tokens", []):
                block(child, "List Number" if token.get("ordered") else "List Bullet")
    elif kind == "table":
        table = doc.add_table(rows=1, cols=len(token["header"]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        count = len(token["header"])
        widths = [1.4, 5.3] if count == 2 else [1.5, 2.6, 2.6] if count == 3 else [6.7 / count] * count
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row_index, values in enumerate([token["header"], *token["rows"]]):
            row = table.rows[0] if row_index == 0 else table.add_row()
            if row_index == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            for index, value in enumerate(values):
                cell = row.cells[index]
                cell.width = Inches(widths[index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ["top", "left", "bottom", "right"]:
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), "80")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                borders = OxmlElement("w:tcBorders")
                for side in ["top", "left", "bottom", "right"]:
                    border = OxmlElement(f"w:{side}")
                    for key, value_attr in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                        border.set(qn(f"w:{key}"), value_attr)
                    borders.append(border)
                props.append(borders)
                shade(cell, "EAF1F3" if row_index == 0 else "F7FAFA" if row_index % 2 == 0 else "FFFFFF")
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                inline(paragraph, value.get("tokens", [{"type": "text", "text": value["text"]}]), bold=row_index == 0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
    elif kind == "code":
        if token.get("diagramPath"):
            image = Path(token["diagramPath"])
            width, height = Image.open(image).size
            physical_width = min(6.7, 3.55 * width / height)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = False
            paragraph.add_run().add_picture(str(image), width=Inches(physical_width))
            picture = paragraph._p.xpath('.//wp:docPr')[0]
            picture.set("descr", "Feature flow diagram generated from the canonical Mermaid source")
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.left_indent = Inches(0.1)
            shade(paragraph, "F1F4F6")
            run = paragraph.add_run(token["text"])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(8.5)
    elif kind == "blockquote":
        for child in token.get("tokens", []):
            block(child)
    elif kind not in ("space", "hr"):
        raise ValueError(f"Unsupported Markdown block: {kind}")


for token in data["tokens"]:
    block(token)
footer = section.footer.paragraphs[0]
footer.paragraph_format.space_after = Pt(0)
footer.add_run(f"Derived from feature.md | {data['hash'][:12]} | ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.save(sys.argv[2])
